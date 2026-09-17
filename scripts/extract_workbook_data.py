from __future__ import annotations

import json
from collections import Counter, defaultdict
from datetime import date, timedelta
from pathlib import Path

import openpyxl
from openpyxl.utils import range_boundaries


ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT / "assets" / "JQS_SOC_CSIRT2.xlsx"
DATA_DIR = ROOT / "data"

TASK_SHEETS = ("Tier I", "Tier II", "Tier III", "Advanced")
INTERNAL_HOST_MARKERS = (
    "dhs.gov",
    "uscis",
    "cisnet",
    "laggar",
    "tines.dhs",
)


def clean(value):
    if value is None:
        return ""
    return str(value).replace("\ufffd", "").strip()


def table_rows(ws, table_name):
    table = ws.tables.get(table_name)
    if table is None:
        return []

    min_col, min_row, max_col, max_row = range_boundaries(table.ref)
    headers = [clean(ws.cell(min_row, col).value) for col in range(min_col, max_col + 1)]
    rows = []

    for row_number in range(min_row + 1, max_row + 1):
        values = [clean(ws.cell(row_number, col).value) for col in range(min_col, max_col + 1)]
        if not any(values):
            continue
        rows.append(dict(zip(headers, values)))

    return rows


def public_url(url):
    if not url:
        return ""
    normalized = url.lower()
    if not normalized.startswith(("http://", "https://")):
        return ""
    if any(marker in normalized for marker in INTERNAL_HOST_MARKERS):
        return ""
    return url


def extract_tasks(wb):
    tasks = []

    for sheet in TASK_SHEETS:
        ws = wb[sheet]
        table_name = "tbl_" + sheet.replace(" ", "_")
        for row in table_rows(ws, table_name):
            task_id = row.get("Task ID", "")
            task = row.get("Duty / Task", "")
            if not task_id or not task:
                continue

            tasks.append(
                {
                    "id": task_id,
                    "title": task,
                    "tier": sheet,
                    "area": row.get("Competency Area", ""),
                    "standard": row.get("Performance Standard", ""),
                    "status": row.get("Status", "Not Started") or "Not Started",
                    "trainer": row.get("Trainer Initials", ""),
                    "notes": row.get("Notes", ""),
                    "owner": row.get("Trainer Initials", "") or "Trainer",
                    "due": "",
                }
            )

    return tasks


def extract_sops(wb):
    rows = []
    for row in table_rows(wb["SOPs"], "tbl_SOPs"):
        name = row.get("ID", "")
        purpose = row.get("Purpose", "")
        if not name and not purpose:
            continue
        rows.append(
            {
                "id": name,
                "title": row.get("Type", "") or name,
                "category": row.get("Type", "General"),
                "status": row.get("Status", ""),
                "owner": "SOC-CSIRT",
                "updated": "",
                "purpose": purpose,
                "sections": [
                    {
                        "heading": "Purpose",
                        "body": purpose or "No SOP purpose was provided in the workbook.",
                    }
                ],
            }
        )
    return rows


def extract_links(wb):
    links = []

    for row in table_rows(wb["Links"], "tbl_Links"):
        name = row.get("Link", "")
        purpose = row.get("Purpose", "")
        url = public_url(row.get("URL", ""))
        if (not name and not purpose) or not url:
            continue
        links.append(
            {
                "name": name or row.get("Type", "Workbook Link"),
                "type": row.get("Type", "Workbook"),
                "category": categorize_link(name or row.get("Type", "Workbook Link")),
                "url": url,
                "description": purpose or f"Reference link for {name or row.get('Type', 'workbook')} workflows.",
            }
        )

    ws = wb["Forensics"]
    for row in ws.iter_rows(min_row=2, values_only=True):
        name = clean(row[0] if len(row) > 0 else "")
        url = public_url(clean(row[1] if len(row) > 1 else ""))
        if not name or not url:
            continue
        links.append(
            {
                "name": name,
                "type": "Forensics",
                "category": categorize_link(name),
                "url": url,
                "description": f"Reference link for {name} workflows.",
            }
        )

    return links


def categorize_link(name):
    normalized = name.lower()
    rules = [
        ("malware", "Malware Analysis"),
        ("cuckoo", "Malware Analysis"),
        ("sandbox", "Malware Analysis"),
        ("volatility", "Memory Forensics"),
        ("memory", "Memory Forensics"),
        ("wireshark", "Network Analysis"),
        ("ncat", "Network Analysis"),
        ("netcat", "Network Analysis"),
        ("hash", "Hashing"),
        ("md5", "Hashing"),
        ("sha", "Hashing"),
        ("registry", "Windows Artifacts"),
        ("reg", "Windows Artifacts"),
        ("prefetch", "Windows Artifacts"),
        ("jumplist", "Windows Artifacts"),
        ("browser", "Browser Artifacts"),
        ("linux", "Linux References"),
        ("unix", "Linux References"),
        ("ubuntu", "Linux References"),
        ("shell", "Linux References"),
        ("sleuth", "Forensic Suites"),
        ("autopsy", "Forensic Suites"),
        ("os forensics", "Forensic Suites"),
        ("ftk", "Forensic Imaging"),
        (" dd ", "Forensic Imaging"),
        ("osfmount", "Forensic Imaging"),
        ("winhex", "Forensic Imaging"),
        ("virtualbox", "Lab Platforms"),
        ("kali", "Lab Platforms"),
        ("reverse", "Reverse Engineering"),
        ("challenges", "Reverse Engineering"),
        ("book", "Reverse Engineering"),
        ("virustotal", "Threat Intelligence"),
        ("sysinternals", "Windows Tools"),
        ("process", "Windows Tools"),
    ]
    padded = f" {normalized} "
    for marker, category in rules:
        if marker in padded:
            return category
    return "Utilities"


def extract_qualifications(wb):
    ws = wb["Quals"]
    quals = []
    current_section = ""

    for row in ws.iter_rows(min_row=1, values_only=True):
        item = clean(row[0] if len(row) > 0 else "")
        if not item:
            continue

        markers = [clean(value) for value in row[1:4]]
        is_header = any(marker in {"Start", "Completed", "Trainer", "Completion Date"} for marker in markers)

        if is_header or item.isupper():
            current_section = item
            continue

        if item.endswith(":"):
            quals.append(
                {
                    "name": item.rstrip(":"),
                    "level": "Signoff",
                    "status": "Pending",
                    "expires": "N/A",
                    "notes": "Workbook qualification signoff field.",
                }
            )
            continue

        quals.append(
            {
                "name": item,
                "level": current_section or "Qualification",
                "status": "Not Started",
                "expires": "N/A",
                "notes": "Imported from the workbook Quals tab.",
            }
        )

    return quals


def build_knowledge(tasks, sops, links):
    by_tier = Counter(task["tier"] for task in tasks)
    by_area = Counter(task["area"] for task in tasks if task["area"])
    tier_articles = [
        {
            "title": f"{tier} Qualification Scope",
            "summary": f"{count} workbook tasks imported for {tier}. Use the Training tab to review task IDs, competency areas, standards, and status.",
            "tags": [tier, "qualification", "training"],
        }
        for tier, count in by_tier.items()
    ]

    top_areas = [
        {
            "title": area,
            "summary": f"{count} JQS tasks are associated with {area}. These tasks appear across the imported tier qualification tables.",
            "tags": [area],
        }
        for area, count in by_area.most_common(8)
    ]

    return [
        {
            "title": "Imported JQS Workbook",
            "summary": f"Loaded {len(tasks)} tiered qualification tasks, {len(sops)} SOP entries, and {len(links)} resource links from JQS_SOC_CSIRT2.xlsx.",
            "tags": ["workbook", "jqs", "overview"],
        },
        *tier_articles,
        *top_areas,
    ]


def build_mock_on_call():
    people = ["Maya Chen", "Andre Patel", "Nina Brooks", "Luis Romero", "Jordan Ellis"]
    pto_rotation = [["Sam Rivera"], ["Taylor Morgan", "Chris Lee"], [], ["Avery Scott"], ["Morgan Blake"], [], ["Riley Park"]]
    start = date(2026, 9, 17)
    schedule = []

    for offset in range(365):
        day = start + timedelta(days=offset)
        primary = people[offset % len(people)]
        backup = people[(offset + 1) % len(people)]
        schedule.append(
            {
                "date": day.isoformat(),
                "primary": primary,
                "backup": backup,
                "pto": pto_rotation[offset % len(pto_rotation)],
            }
        )

    return schedule


def write_json(filename, data):
    DATA_DIR.mkdir(exist_ok=True)
    path = DATA_DIR / filename
    path.write_text(json.dumps(data, indent=2) + "\n", encoding="utf-8")


def main():
    if not WORKBOOK.exists():
        raise SystemExit(f"Workbook not found: {WORKBOOK}")

    wb = openpyxl.load_workbook(WORKBOOK, data_only=True)

    tasks = extract_tasks(wb)
    sops = extract_sops(wb)
    links = extract_links(wb)
    knowledge = build_knowledge(tasks, sops, links)

    write_json("training.json", tasks)
    write_json("sops.json", sops)
    write_json("links.json", links)
    write_json("knowledge.json", knowledge)
    write_json("on-call.json", build_mock_on_call())

    print(f"Imported {len(tasks)} training tasks")
    print(f"Imported {len(sops)} SOP entries")
    print(f"Imported {len(links)} links")
    print("Skipped certs.json; Certs is maintained as a cybersecurity certification catalog")
    print(f"Generated {len(knowledge)} knowledge summaries")
    print("Generated 365 mock on-call schedule records")


if __name__ == "__main__":
    main()
